#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
MDs - Ultimate Markdown Converter

Mengonversi dokumen/kode ke Markdown, membundel folder proyek menjadi satu
file .md, dan memulihkannya kembali ke file asli.

Mode interaktif : MDc            (atau: python md.py)
Seret & lepas   : MDc <file/folder>  (atau seret ke MDc.bat di Windows)
Mode CLI        : MDc file    <path>   [-o OUT] [--embed] [--force]
                  MDc bundle  <folder> [-o OUT] [--ext .py,.js] [--ignore POLA]
                                       [--max-kb N] [--dry-run] [--force]
                  MDc restore <file.md> [-d TUJUAN] [--name NAMA] [--text]
                                       [--rebuild] [--dry-run] [--force]
                  MDc doctor           (periksa instalasi: Python, pustaka, PATH, git)
"""

import argparse
import base64
import fnmatch
import hashlib
import os
import re
import shutil
import subprocess
import sys
import unicodedata
import zlib
from datetime import datetime
from urllib.parse import unquote, urlparse

# Pustaka opsional: skrip tetap jalan (untuk file teks/kode) bila belum terpasang.
try:
    import pypdf
except ImportError:
    pypdf = None

try:
    from docx import Document as DocxDocument
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph
except ImportError:
    DocxDocument = None

try:  # python-docx >= 1.0
    from docx.text.hyperlink import Hyperlink as DocxHyperlink
except ImportError:
    DocxHyperlink = None

VERSION = "2.2"

# --------------------------------------------------------------------------
# Konfigurasi
# --------------------------------------------------------------------------

# Ekstensi teks/kode -> tag syntax highlighting Markdown.
LANG_MAP = {
    # teks & dokumentasi
    ".txt": "text", ".md": "markdown", ".markdown": "markdown", ".mdx": "mdx",
    ".rst": "rst", ".adoc": "asciidoc", ".tex": "latex", ".bib": "bibtex",
    ".org": "org", ".csv": "csv", ".tsv": "tsv", ".diff": "diff", ".patch": "diff",
    # konfigurasi & data
    ".json": "json", ".jsonc": "jsonc", ".json5": "json5",
    ".yaml": "yaml", ".yml": "yaml", ".toml": "toml", ".ini": "ini",
    ".cfg": "ini", ".conf": "conf", ".properties": "properties",
    ".xml": "xml", ".plist": "xml", ".csproj": "xml", ".service": "ini",
    ".tf": "hcl", ".tfvars": "hcl", ".nix": "nix", ".proto": "protobuf",
    ".graphql": "graphql", ".gql": "graphql", ".cmake": "cmake",
    # web
    ".html": "html", ".htm": "html", ".xhtml": "html", ".css": "css",
    ".scss": "scss", ".sass": "sass", ".less": "less",
    ".js": "javascript", ".mjs": "javascript", ".cjs": "javascript",
    ".jsx": "jsx", ".ts": "typescript", ".tsx": "tsx",
    ".vue": "vue", ".svelte": "svelte", ".php": "php",
    # bahasa pemrograman
    ".py": "python", ".c": "c", ".h": "c", ".cpp": "cpp", ".cc": "cpp",
    ".cxx": "cpp", ".hpp": "cpp", ".hh": "cpp", ".hxx": "cpp", ".ino": "cpp",
    ".java": "java", ".cs": "csharp", ".go": "go", ".rs": "rust",
    ".rb": "ruby", ".kt": "kotlin", ".kts": "kotlin", ".swift": "swift",
    ".dart": "dart", ".scala": "scala", ".lua": "lua", ".r": "r",
    ".pl": "perl", ".pm": "perl", ".hs": "haskell", ".ex": "elixir",
    ".exs": "elixir", ".erl": "erlang", ".clj": "clojure", ".jl": "julia",
    ".zig": "zig", ".sol": "solidity", ".gd": "gdscript", ".vb": "vbnet",
    ".fs": "fsharp", ".groovy": "groovy", ".gradle": "groovy", ".sql": "sql",
    ".asm": "asm",
    # shell & skrip
    ".sh": "bash", ".bash": "bash", ".zsh": "zsh", ".fish": "fish",
    ".bat": "batch", ".cmd": "batch", ".ps1": "powershell", ".psm1": "powershell",
}
# File tanpa ekstensi yang umum (dicocokkan lewat nama, huruf kecil).
SPECIAL_NAMES = {
    "dockerfile": "dockerfile", "makefile": "makefile", "gnumakefile": "makefile",
    "license": "text", "licence": "text", "copying": "text", "notice": "text",
    "authors": "text", "changelog": "text", "readme": "text", "procfile": "text",
    "gemfile": "ruby", "rakefile": "ruby", "vagrantfile": "ruby",
    "jenkinsfile": "groovy", "brewfile": "ruby",
}
# File tersembunyi yang aman & berguna (".env" dan sejenisnya TIDAK termasuk).
SAFE_DOTFILES = {
    ".gitignore": "gitignore", ".gitattributes": "gitattributes",
    ".dockerignore": "gitignore", ".editorconfig": "ini", ".prettierrc": "json",
    ".eslintrc": "json", ".babelrc": "json", ".nvmrc": "text",
    ".python-version": "text", ".flake8": "ini",
}
TEXT_EXTS = set(LANG_MAP)
DOC_EXTS = {".pdf", ".docx"}
SUPPORTED_EXTS = TEXT_EXTS | DOC_EXTS
MD_EXTS = {".md", ".markdown"}


def lang_for(name):
    """Tag bahasa untuk nama file teks yang didukung, atau None bila tidak didukung."""
    low = name.lower()
    ext = os.path.splitext(low)[1]
    if ext in LANG_MAP:
        return LANG_MAP[ext]
    return SPECIAL_NAMES.get(low) or SAFE_DOTFILES.get(low)


# Folder & file yang dilewati saat bundle (noise / besar / berisiko).
IGNORE_DIRS = {
    "node_modules", "venv", "env", "__pycache__", "dist", "build",
    "target", "vendor", "site-packages",
}
IGNORE_FILES = {
    "package-lock.json", "yarn.lock", "pnpm-lock.yaml", "poetry.lock",
    "composer.lock", "Cargo.lock",
}
MAX_FILE_BYTES = 1_000_000  # file teks di atas ini dilewati saat bundle
PAGE_SIZE = 20             # jumlah item per halaman di navigasi folder
BUNDLE_TITLE = "# Bundle Proyek:"

# --------------------------------------------------------------------------
# Tampilan terminal
# --------------------------------------------------------------------------

if os.name == "nt":
    os.system("")  # aktifkan ANSI escape di Windows 10+

RESET = "\033[0m"
RED = "\033[1;31m"
GREEN = "\033[1;32m"
YELLOW = "\033[1;33m"
BLUE = "\033[1;34m"
CYAN = "\033[1;36m"
MAGENTA = "\033[0;35m"

BANNER = r"""
                   /$$      /$$ /$$$$$$$
                  | $$$    /$$$| $$__  $$
                  | $$$$  /$$$$| $$  \ $$  /$$$$$$$
                  | $$ $$/$$ $$| $$  | $$ /$$_____/
                  | $$  $$$| $$| $$  | $$| $$
                  | $$\  $ | $$| $$  | $$| $$
                  | $$ \/  | $$| $$$$$$$/|  $$$$$$$
                  |__/     |__/|_______/  \_______/



  /$$$$$$  /$$       /$$                         /$$     /$$
 /$$__  $$| $$      |__/                        |  $$   /$$/
| $$  \__/| $$$$$$$  /$$ /$$$$$$/$$$$   /$$$$$$  \  $$ /$$/  /$$   /$$ /$$   /$$
|  $$$$$$ | $$__  $$| $$| $$_  $$_  $$ |____  $$  \  $$$$/  | $$  | $$| $$  | $$
 \____  $$| $$  \ $$| $$| $$ \ $$ \ $$  /$$$$$$$   \  $$/   | $$  | $$| $$  | $$
 /$$  \ $$| $$  | $$| $$| $$ | $$ | $$ /$$__  $$    | $$    | $$  | $$| $$  | $$
|  $$$$$$/| $$  | $$| $$| $$ | $$ | $$|  $$$$$$$    | $$    |  $$$$$$/|  $$$$$$/
 \______/ |__/  |__/|__/|__/ |__/ |__/ \_______/    |__/     \______/  \______/
"""


def clear_screen():
    os.system("cls" if os.name == "nt" else "clear")


def print_banner():
    clear_screen()
    width = shutil.get_terminal_size((80, 24)).columns
    if width >= 80:
        art = "\n".join(line.rstrip() for line in BANNER.strip("\n").split("\n"))
        print(CYAN + art + RESET)
    else:  # layar sempit (mis. Termux portrait): banner ringkas
        print(CYAN + "=== MDs | ShimaYuu ===" + RESET)
    print(f"{YELLOW}=== ULTIMATE MARKDOWN CONVERTER v{VERSION} ==={RESET}")
    print(f"{GREEN}Format:{RESET} PDF & DOCX (dua arah), kode/teks 90+ ekstensi, Dockerfile, dll.")
    print("-" * min(width, 50))


def pause(msg="\nTekan [Enter] untuk lanjut..."):
    input(msg)


def ask(prompt):
    """input() yang mengenali kata 'exit' di mana pun."""
    value = input(prompt).strip()
    if value.lower() == "exit":
        print("\nKeluar dari skrip. Sampai jumpa!")
        sys.exit(0)
    return value


def confirm(prompt):
    return ask(f"{prompt} (y/n): ").lower() in ("y", "ya")


def format_size(num_bytes):
    if num_bytes < 1024:
        return f"{num_bytes} B"
    return f"{num_bytes / 1024:.1f} KB"


def estimate_tokens(text):
    """Perkiraan kasar (~4 karakter per token). Blok salinan asli tidak dihitung."""
    cut = text.rfind("<!-- META_ORIGINAL_BASE64")
    if cut >= 0:
        text = text[:cut]
    return max(1, len(text) // 4)


# --------------------------------------------------------------------------
# Utilitas file
# --------------------------------------------------------------------------

def read_text_file(path):
    """Baca file teks dengan fallback encoding; tolak file biner."""
    with open(path, "rb") as f:
        raw = f.read()
    if b"\0" in raw[:8192]:
        raise ValueError("file biner")
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")
    return text.replace("\r\n", "\n").replace("\r", "\n")


def write_text_file(path, text):
    parent = os.path.dirname(path)
    if parent:
        os.makedirs(parent, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)


def sanitize_filename(name):
    """Ambil nama file saja (buang path) dan tolak nama berbahaya."""
    name = os.path.basename(name.replace("\\", "/")).strip()
    return "" if name in ("", ".", "..") else name


def safe_join(base_dir, rel_path):
    """Gabungkan path dan pastikan hasilnya tetap di dalam base_dir."""
    rel_path = rel_path.replace("\\", "/")
    base_real = os.path.realpath(base_dir)
    target = os.path.realpath(os.path.join(base_real, rel_path))
    try:
        if os.path.commonpath([base_real, target]) != base_real:
            return None
    except ValueError:  # beda drive di Windows
        return None
    return target if target != base_real else None


def default_start_dir():
    start = os.path.expanduser("~/storage/downloads")  # Termux
    return start if os.path.isdir(start) else os.getcwd()


def copy_to_clipboard(text):
    """Salin ke clipboard (hanya Termux). Return True bila berhasil."""
    exe = shutil.which("termux-clipboard-set")
    if not exe:
        return False
    try:
        subprocess.run([exe], input=text.encode("utf-8"), check=True, timeout=15)
        return True
    except Exception:
        return False


# --------------------------------------------------------------------------
# Ekstraksi dokumen (PDF / DOCX)
# --------------------------------------------------------------------------

def pdf_to_text(path):
    if pypdf is None:
        raise ValueError("pypdf belum terpasang (jalankan: python -m pip install pypdf)")
    reader = pypdf.PdfReader(path)
    if reader.is_encrypted:
        try:
            ok = reader.decrypt("")
        except Exception:
            ok = 0
        if not ok:
            raise ValueError("PDF terenkripsi / butuh password")
    parts = []
    for number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(f"<!-- Halaman {number} -->\n\n{text}")
    if not parts:
        raise ValueError("tidak ada teks yang bisa diekstrak (mungkin PDF hasil scan, butuh OCR)")
    return "\n\n".join(parts) + "\n"


# ---------------------------------------------------------------------------
# Escape Markdown (supaya teks dokumen tidak salah terbaca saat dipulihkan)
# ---------------------------------------------------------------------------

def _esc_inline(text):
    return re.sub(r"([\\`*\[\]|])", r"\\\1", text)


def _esc_block(text):
    """Escape penanda blok di awal paragraf (#, -, +, >, 1.)."""
    match = re.match(r"^(\d+)([.)])(\s|$)", text)
    if match:
        return match.group(1) + "\\" + text[match.end(1):]
    if text[:1] in ("#", "-", "+", ">"):
        return "\\" + text
    return text


def _safe_url(url):
    if not url or not re.match(r"(?i)^(https?://|mailto:)", url):
        return None
    return url.replace(" ", "%20").replace(")", "%29").replace("(", "%28")


def _docx_inline(par):
    """Teks paragraf dengan **tebal**, *miring*, dan [tautan](url)."""
    items = []  # [teks, tebal, miring, url]

    def add(text, bold, italic, url=None):
        if text:
            items.append([text.replace("\r", ""), bool(bold), bool(italic), url])

    if hasattr(par, "iter_inner_content"):
        for item in par.iter_inner_content():
            if DocxHyperlink is not None and isinstance(item, DocxHyperlink):
                runs = list(item.runs)
                add(item.text, runs and all(r.bold for r in runs),
                    runs and all(r.italic for r in runs), _safe_url(item.url))
            else:
                add(item.text, item.bold, item.italic)
    else:
        add(par.text, False, False)

    merged = []
    for item in items:
        if merged and merged[-1][1:] == item[1:]:
            merged[-1][0] += item[0]
        else:
            merged.append(item)

    out = []
    for text, bold, italic, url in merged:
        core = text.strip()
        if not core:
            out.append(text.replace("\n", "  \n"))
            continue
        lead = text[:len(text) - len(text.lstrip())].replace("\n", "  \n")
        trail = text[len(text.rstrip()):].replace("\n", "  \n")
        core_md = _esc_inline(core).replace("\n", "  \n")
        if url:
            core_md = f"[{core_md}]({url})"
        mark = "***" if bold and italic else "**" if bold else "*" if italic else ""
        out.append(f"{lead}{mark}{core_md}{mark}{trail}")
    return "".join(out).strip()


def _docx_list_info(par, doc):
    """Deteksi daftar bawaan Word (numPr). Return (level, berurut) atau None."""
    pPr = par._p.pPr
    num_pr = pPr.numPr if pPr is not None else None
    if num_pr is None or num_pr.numId is None or num_pr.numId.val == 0:
        return None
    level = int(num_pr.ilvl.val) if num_pr.ilvl is not None else 0
    ordered = False
    try:
        numbering = doc.part.numbering_part.element
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        w = "{%s}" % ns["w"]
        num_id = str(num_pr.numId.val)
        for num in numbering.findall("w:num", ns):
            if num.get(w + "numId") != num_id:
                continue
            abs_id = num.find("w:abstractNumId", ns).get(w + "val")
            for absn in numbering.findall("w:abstractNum", ns):
                if absn.get(w + "abstractNumId") != abs_id:
                    continue
                for lvl in absn.findall("w:lvl", ns):
                    if lvl.get(w + "ilvl") == str(level):
                        fmt = lvl.find("w:numFmt", ns)
                        ordered = fmt is not None and fmt.get(w + "val") != "bullet"
    except Exception:
        pass
    return level, ordered


def _docx_paragraph(par, doc=None):
    if not par.text.strip():
        return ("p", "")
    style = (par.style.name if par.style is not None else "") or ""
    low = style.lower()
    if low == "title":
        return ("h", "# " + _esc_inline(par.text.strip()))
    match = re.match(r"heading (\d)", low)
    if match:
        return ("h", "#" * min(int(match.group(1)), 6) + " " + _esc_inline(par.text.strip()))

    level, ordered, is_list = 0, False, False
    style_match = re.match(r"list (bullet|number)(?: (\d))?$", low)
    if style_match:
        is_list = True
        ordered = style_match.group(1) == "number"
        level = int(style_match.group(2) or 1) - 1
    elif doc is not None:
        info = _docx_list_info(par, doc)
        if info:
            is_list = True
            level, ordered = info
    text = _docx_inline(par)
    if is_list:
        return ("li", "    " * level + ("1. " if ordered else "- ") + text)
    return ("p", _esc_block(text))


def _docx_table(table):
    rows = []
    for row in table.rows:
        cells, seen = [], set()
        for c in row.cells:
            dup = id(c._tc) in seen  # sel gabungan muncul berulang
            seen.add(id(c._tc))
            cells.append("" if dup else _esc_inline(c.text.strip()).replace("\n", " "))
        rows.append(cells)
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def docx_to_markdown(path):
    if DocxDocument is None:
        raise ValueError("python-docx belum terpasang (jalankan: python -m pip install python-docx)")
    doc = DocxDocument(path)
    blocks = []  # (jenis, teks)
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            kind, text = _docx_paragraph(DocxParagraph(child, doc), doc)
        elif tag == "tbl":
            kind, text = "t", _docx_table(DocxTable(child, doc))
        else:
            continue
        if text.strip():
            blocks.append((kind, text))

    out = ""
    prev = None
    for kind, text in blocks:
        if prev is not None:
            out += "\n" if (kind == "li" and prev == "li") else "\n\n"
        out += text
        prev = kind
    return out + "\n"


# --------------------------------------------------------------------------
# Markdown -> DOCX (membangun ulang dokumen Word dari teks Markdown)
# --------------------------------------------------------------------------

_MD_ESCAPABLE = r"\\`*_\[\]()#+\-.!|>{}~"
_INLINE_RE = re.compile(
    r"(?P<esc>\\[" + _MD_ESCAPABLE + r"])"
    r"|(?P<code>`(?P<ct>[^`\n]+)`)"
    r"|(?P<link>\[(?P<lt>(?:\\.|[^\]\\])+)\]\((?P<lu>[^)\s]+)\))"
    r"|(?P<bi>\*\*\*(?=[^\s*])(?P<bit>.+?)(?<=[^\s*])\*\*\*)"
    r"|(?P<b>\*\*(?=[^\s*])(?P<bt>.+?)(?<=[^\s*])\*\*)"
    r"|(?P<i>\*(?=[^\s*])(?P<it>.+?)(?<=[^\s*])\*)",
    re.S)


def _parse_inline(text, bold=False, italic=False, url=None):
    """Return [(teks, tebal, miring, kode, url)]."""
    out, pos = [], 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], bold, italic, False, url))
        if m.group("esc"):
            out.append((m.group("esc")[1], bold, italic, False, url))
        elif m.group("code"):
            out.append((m.group("ct"), bold, italic, True, url))
        elif m.group("link"):
            out.extend(_parse_inline(m.group("lt"), bold, italic, _safe_url(m.group("lu"))))
        elif m.group("bi"):
            out.extend(_parse_inline(m.group("bit"), True, True, url))
        elif m.group("b"):
            out.extend(_parse_inline(m.group("bt"), True, italic, url))
        else:
            out.extend(_parse_inline(m.group("it"), bold, True, url))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], bold, italic, False, url))
    return out


def _docx_add_hyperlink(par, text, url, bold, italic):
    r_id = par.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    if bold:
        props.append(OxmlElement("w:b"))
    if italic:
        props.append(OxmlElement("w:i"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    link.append(run)
    par._p.append(link)


def _docx_add_inline(par, text, force_bold=False):
    for seg, bold, italic, code, url in _parse_inline(text):
        bold = bold or force_bold
        if url:
            _docx_add_hyperlink(par, seg, url, bold, italic)
            continue
        run = par.add_run(seg)
        run.bold = bold or None
        run.italic = italic or None
        if code:
            run.font.name = "Courier New"


def _docx_restart_numbering(doc):
    """Buat instance daftar berurut baru yang mulai dari 1. Return numId atau None."""
    try:
        numbering = doc.part.numbering_part.element
        style_num = doc.styles["List Number"].element.pPr.numPr.numId.val
        abstract = None
        for num in numbering.findall(qn("w:num")):
            if int(num.get(qn("w:numId"))) == style_num:
                abstract = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        if abstract is None:
            return None
        new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(new_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), abstract)
        num.append(ref)
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:startOverride")
        start.set(qn("w:val"), "1")
        override.append(start)
        num.append(override)
        numbering.append(num)
        return new_id
    except Exception:
        return None


def _docx_rule(doc):
    par = doc.add_paragraph()
    props = par._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, val in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "999999")):
        bottom.set(qn("w:" + key), val)
    border.append(bottom)
    props.append(border)


PAGE_MARK_RE = re.compile(r"^<!--\s*Halaman\s+\d+\s*-->\s*$", re.M)
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
_FENCE_RE = re.compile(r"^\s*(`{3,}|~{3,})")
_LIST_RE = re.compile(r"^(\s*)(?:([-*+])|(\d+)[.)])\s+(.*)$")
_HEAD_RE = re.compile(r"^\s{0,3}(#{1,6})\s+(.*?)\s*#*\s*$")
_HR_RE = re.compile(r"^\s{0,3}([-*_])(\s*\1){2,}\s*$")


def _split_table_row(row):
    row = row.strip()
    cells = re.split(r"(?<!\\)\|", row)
    if row.startswith("|"):
        cells = cells[1:]
    if row.endswith("|") and not row.endswith("\\|") and cells:
        cells = cells[:-1]
    return [c.strip() for c in cells]


def _md_blocks(md_text):
    """Pecah Markdown menjadi blok: (jenis, data). Dipakai oleh DOCX dan PDF."""
    lines = md_text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks, para, i = [], [], 0
    stack = []  # tumpukan indentasi daftar

    def flush():
        if para:
            blocks.append(("p", list(para)))
            para.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            flush()
            i += 1
            continue
        if PAGE_MARK_RE.match(line):
            flush()
            stack.clear()
            blocks.append(("pagebreak", None))
            i += 1
            continue
        if re.match(r"^\s*<!--.*-->\s*$", line):
            i += 1
            continue

        fence = _FENCE_RE.match(line)
        if fence:
            flush()
            stack.clear()
            mark, code = fence.group(1), []
            i += 1
            while i < len(lines):
                closing = lines[i].strip()
                if closing.startswith(mark[0] * len(mark)) and set(closing) == {mark[0]}:
                    break
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", code))
            continue

        head = _HEAD_RE.match(line)
        if head:
            flush()
            stack.clear()
            blocks.append(("h", (len(head.group(1)), head.group(2))))
            i += 1
            continue

        if _HR_RE.match(line):
            flush()
            stack.clear()
            blocks.append(("hr", None))
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and _TABLE_SEP_RE.match(lines[i + 1]):
            flush()
            stack.clear()
            rows = [_split_table_row(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_table_row(lines[i]))
                i += 1
            blocks.append(("table", rows))
            continue

        if stripped.startswith(">"):
            flush()
            stack.clear()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            blocks.append(("quote", quote))
            continue

        item = _LIST_RE.match(line)
        if item:
            flush()
            indent = len(item.group(1).replace("\t", "    "))
            while stack and indent < stack[-1]:
                stack.pop()
            if not stack or indent > stack[-1]:
                stack.append(indent)
            text = item.group(4)
            i += 1
            while (i < len(lines) and lines[i].strip() and lines[i][:1] in (" ", "\t")
                   and not _LIST_RE.match(lines[i])):
                text += " " + lines[i].strip()
                i += 1
            blocks.append(("li", (len(stack) - 1, item.group(3) is not None, text)))
            continue

        stack.clear()
        para.append(line)
        i += 1
    flush()
    return blocks


def _join_paragraph(lines, keep_breaks):
    """Gabungkan baris paragraf. Akhiran dua spasi / '\\' = pindah baris."""
    if keep_breaks:
        return "\n".join(l.rstrip() for l in lines)
    out = ""
    for index, line in enumerate(lines):
        hard = line.endswith("  ") or (line.endswith("\\") and not line.endswith("\\\\"))
        text = line.strip()
        if line.endswith("\\") and not line.endswith("\\\\"):
            text = text[:-1].rstrip()
        out += text
        if index < len(lines) - 1:
            out += "\n" if hard else " "
    return out


def md_to_docx(md_text, out_path, title=None):
    """Bangun file .docx dari teks Markdown. Return None."""
    if DocxDocument is None:
        raise ValueError("python-docx belum terpasang (jalankan: python -m pip install python-docx)")
    plain = bool(PAGE_MARK_RE.search(md_text))  # hasil ekstraksi PDF: teks apa adanya
    doc = DocxDocument()
    if title:
        doc.core_properties.title = title
    num_id = None
    seen_page = False
    first_block = True

    if plain:
        for kind, data in _md_blocks_plain(md_text):
            if kind == "pagebreak":
                if not first_block:
                    doc.add_page_break()
                continue
            par = doc.add_paragraph()
            par.add_run(data)
            first_block = False
        doc.save(out_path)
        return

    for kind, data in _md_blocks(md_text):
        if kind != "li":
            num_id = None
        if kind == "pagebreak":
            if seen_page:
                doc.add_page_break()
            seen_page = True
            continue
        if kind == "h":
            level, text = data
            par = doc.add_heading("", level=min(level, 9))
            _docx_add_inline(par, text)
        elif kind == "p":
            par = doc.add_paragraph()
            _docx_add_inline(par, _join_paragraph(data, False))
        elif kind == "li":
            level, ordered, text = data
            level = min(level, 2)
            base = "List Number" if ordered else "List Bullet"
            style = base if level == 0 else f"{base} {level + 1}"
            par = doc.add_paragraph(style=style)
            if ordered and level == 0:
                if num_id is None:
                    num_id = _docx_restart_numbering(doc)
                if num_id is not None:
                    num_pr = par._p.get_or_add_pPr().get_or_add_numPr()
                    num_pr.get_or_add_ilvl().val = 0
                    num_pr.get_or_add_numId().val = num_id
            _docx_add_inline(par, text)
        elif kind == "code":
            for code_line in data or [""]:
                par = doc.add_paragraph()
                par.paragraph_format.space_after = Pt(0)
                run = par.add_run(code_line.replace("\t", "    "))
                run.font.name = "Courier New"
                run.font.size = Pt(9)
        elif kind == "quote":
            par = doc.add_paragraph(style="Quote")
            _docx_add_inline(par, _join_paragraph(data, False))
        elif kind == "hr":
            _docx_rule(doc)
        elif kind == "table":
            width = max(len(r) for r in data)
            table = doc.add_table(rows=len(data), cols=width)
            table.style = "Table Grid"
            for r, row in enumerate(data):
                for c in range(width):
                    cell = table.cell(r, c)
                    _docx_add_inline(cell.paragraphs[0], row[c] if c < len(row) else "",
                                     force_bold=(r == 0))
            doc.add_paragraph()
    doc.save(out_path)


def _md_blocks_plain(text):
    """Mode teks apa adanya (hasil ekstraksi PDF): paragraf = blok antar baris kosong."""
    out, buf = [], []
    for line in text.replace("\r\n", "\n").split("\n"):
        if PAGE_MARK_RE.match(line):
            if buf:
                out.append(("p", "\n".join(buf)))
                buf = []
            out.append(("pagebreak", None))
        elif re.match(r"^\s*<!--.*-->\s*$", line):
            continue
        elif not line.strip():
            if buf:
                out.append(("p", "\n".join(buf)))
                buf = []
        else:
            buf.append(line.rstrip())
    if buf:
        out.append(("p", "\n".join(buf)))
    return out


# --------------------------------------------------------------------------
# Markdown -> PDF (penulis PDF mini tanpa dependensi; font bawaan PDF)
# --------------------------------------------------------------------------

# Lebar glyph Helvetica (per 1000 unit) untuk ASCII 32..126.
_HELV_W = [
    278, 278, 355, 556, 556, 889, 667, 191, 333, 333, 389, 584, 278, 333, 278, 278,
    556, 556, 556, 556, 556, 556, 556, 556, 556, 556, 278, 278, 584, 584, 584, 556,
    1015, 667, 667, 722, 722, 667, 611, 778, 722, 278, 500, 667, 556, 833, 722, 778,
    667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611, 278, 278, 278, 469, 556,
    333, 556, 556, 500, 556, 556, 278, 556, 556, 222, 222, 500, 222, 833, 556, 556,
    556, 556, 333, 500, 278, 556, 500, 722, 500, 500, 500, 334, 260, 334, 584]
_HELV_BOLD_W = [
    278, 333, 474, 556, 556, 889, 722, 238, 333, 333, 389, 584, 278, 333, 278, 278,
    556, 556, 556, 556, 556, 556, 556, 556, 556, 556, 333, 333, 584, 584, 584, 611,
    975, 722, 722, 722, 722, 667, 611, 778, 722, 278, 556, 722, 611, 833, 722, 778,
    667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611, 333, 278, 333, 584, 556,
    333, 556, 611, 556, 611, 556, 333, 611, 611, 278, 278, 556, 278, 889, 611, 611,
    611, 611, 389, 556, 333, 611, 556, 778, 556, 556, 500, 389, 280, 389, 584]
_SPECIAL_W = {  # (reguler, tebal)
    "•": (350, 350), "–": (556, 556), "—": (1000, 1000), "‘": (222, 278),
    "’": (222, 278), "“": (333, 500), "”": (333, 500), "…": (1000, 1000),
    "€": (556, 556), "©": (737, 737), "®": (737, 737), "°": (400, 400),
    "×": (584, 584), "÷": (584, 584), "±": (584, 584), "«": (556, 556),
    "»": (556, 556), "§": (556, 556), "ß": (611, 611), "æ": (889, 889),
    "Æ": (1000, 1000), "œ": (944, 944), "Œ": (1000, 1000), "ø": (611, 611),
    "Ø": (778, 778), "™": (1000, 1000), "½": (834, 834), "·": (278, 278),
}
_PDF_SUBST = {
    "→": "->", "←": "<-", "⇒": "=>", "✓": "v", "✔": "v", "✗": "x", "✘": "x",
    "≥": ">=", "≤": "<=", "≠": "!=", "−": "-", "‐": "-", "‑": "-", "‒": "-",
    "―": "-", "\u00a0": " ", "\u2009": " ", "\u202f": " ", "\u200b": "",
    "\u200d": "", "\ufe0f": "", "\u2028": "\n", "★": "*", "☆": "*", "▪": "-",
    "▸": ">", "►": ">", "▶": ">", "✅": "[v]", "❌": "[x]",
}
_PDF_FONTS = {  # kunci -> (nama resource, BaseFont)
    "r": ("F1", "Helvetica"), "b": ("F2", "Helvetica-Bold"),
    "i": ("F3", "Helvetica-Oblique"), "bi": ("F4", "Helvetica-BoldOblique"),
    "c": ("F5", "Courier"),
}


def _font_key(bold=False, italic=False, code=False):
    if code:
        return "c"
    return "bi" if bold and italic else "b" if bold else "i" if italic else "r"


def _char_w(ch, fk):
    if fk == "c":
        return 600
    bold = fk in ("b", "bi")
    table = _HELV_BOLD_W if bold else _HELV_W
    code = ord(ch)
    if 32 <= code <= 126:
        return table[code - 32]
    special = _SPECIAL_W.get(ch)
    if special:
        return special[1 if bold else 0]
    base = unicodedata.normalize("NFD", ch)[0]
    if 32 <= ord(base) <= 126:
        return table[ord(base) - 32]
    return 556


def _text_w(text, fk, size):
    return sum(_char_w(c, fk) for c in text) * size / 1000.0


def _to_winansi(text, lost):
    """Ubah teks agar aman untuk font bawaan PDF (WinAnsi). Hitung karakter yang hilang."""
    out = []
    for ch in text:
        if ch == "\n":
            out.append(ch)
        elif ch == "\t":
            out.append("    ")
        elif ch in _PDF_SUBST:
            out.append(_PDF_SUBST[ch])
        else:
            try:
                ch.encode("cp1252")
                if ch >= " ":
                    out.append(ch)
                continue
            except UnicodeEncodeError:
                pass
            if unicodedata.category(ch) in ("Mn", "Cf", "Cc", "Me"):
                continue
            code = ord(ch)
            if code >= 0x1F000 or 0x2190 <= code <= 0x2BFF:  # emoji & simbol dekoratif
                lost[1] += 1
                continue
            ascii_form = "".join(c for c in unicodedata.normalize("NFKD", ch) if ord(c) < 128)
            if ascii_form:
                out.append(ascii_form)
            else:
                out.append("?")
                lost[0] += 1
    return "".join(out)


def _pdf_str(text):
    raw = text.encode("cp1252", "replace").decode("latin-1")
    return raw.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


class _PdfBuilder:
    W, H, MARGIN = 595.28, 841.89, 56.0

    def __init__(self):
        self.pages = []
        self.lost = [0, 0]  # [karakter diganti '?', emoji/simbol dihapus]
        self.new_page()

    @property
    def content_w(self):
        return self.W - 2 * self.MARGIN

    def new_page(self):
        self.page = {"ops": [], "links": []}
        self.pages.append(self.page)
        self.y = self.H - self.MARGIN

    def at_top(self):
        return self.y >= self.H - self.MARGIN - 0.5

    def ensure(self, height):
        if self.y - height < self.MARGIN and not self.at_top():
            self.new_page()

    def page_break(self):
        if not self.at_top():
            self.new_page()

    # -- teks ---------------------------------------------------------------
    def segments(self, text, base_bold=False):
        out = []
        for seg, bold, italic, code, url in _parse_inline(text):
            seg = _to_winansi(seg, self.lost)
            if seg:
                out.append((seg, _font_key(bold or base_bold, italic, code), url))
        return out

    def wrap(self, segments, size, width):
        lines, cur = [[]], 0.0
        for text, fk, url in segments:
            for tok in re.findall(r"\n|[^\S\n]+|\S+", text):
                if tok == "\n":
                    lines.append([])
                    cur = 0.0
                    continue
                w = _text_w(tok, fk, size)
                if tok.isspace():
                    if lines[-1]:
                        lines[-1].append((tok, fk, url))
                        cur += w
                    continue
                if cur + w > width and lines[-1]:
                    while lines[-1] and lines[-1][-1][0].isspace():
                        lines[-1].pop()
                    lines.append([])
                    cur = 0.0
                if w > width:  # kata terlalu panjang: pecah per karakter
                    piece = ""
                    for ch in tok:
                        if _text_w(piece + ch, fk, size) > width and piece:
                            lines[-1].append((piece, fk, url))
                            lines.append([])
                            piece = ch
                        else:
                            piece += ch
                    lines[-1].append((piece, fk, url))
                    cur = _text_w(piece, fk, size)
                else:
                    lines[-1].append((tok, fk, url))
                    cur += w
        for line in lines:
            while line and line[-1][0].isspace():
                line.pop()
        return lines

    def emit_line(self, x, baseline, tokens, size):
        parts = ["BT", f"{x:.2f} {baseline:.2f} Td"]
        font, blue, cur = None, False, x
        for text, fk, url in tokens:
            if not text:
                continue
            if fk != font:
                parts.append(f"/{_PDF_FONTS[fk][0]} {size:.2f} Tf")
                font = fk
            if bool(url) != blue:
                parts.append("0.02 0.27 0.8 rg" if url else "0 g")
                blue = bool(url)
            parts.append(f"({_pdf_str(text)}) Tj")
            width = _text_w(text, fk, size)
            if url:
                self.page["links"].append(
                    (cur, baseline - size * 0.2, cur + width, baseline + size * 0.85, url))
            cur += width
        parts += ["0 g", "ET"]
        self.page["ops"].append("\n".join(parts))

    def paragraph(self, segments, size=11.0, indent=0.0, marker=None, marker_x=None,
                  after=6.0, lead=1.4, bar=False):
        lines = self.wrap(segments, size, self.content_w - indent)
        lh = size * lead
        for index, line in enumerate(lines):
            self.ensure(lh)
            baseline = self.y - size
            if bar:
                self.page["ops"].append(
                    f"0.75 g {self.MARGIN + indent - 8:.2f} {self.y - lh:.2f} 2.5 {lh:.2f} re f 0 g")
            if index == 0 and marker:
                mx = self.MARGIN + (marker_x if marker_x is not None else indent - 14)
                self.emit_line(mx, baseline, [(marker, "r", None)], size)
            if line:
                self.emit_line(self.MARGIN + indent, baseline, line, size)
            self.y -= lh
        self.y -= after

    def code(self, code_lines, size=9.0):
        per_line = max(10, int((self.content_w - 12) / (0.6 * size)))
        lh = size * 1.3
        for raw in code_lines or [""]:
            text = _to_winansi(raw, self.lost)
            chunks = [text[i:i + per_line] for i in range(0, len(text), per_line)] or [""]
            for chunk in chunks:
                self.ensure(lh)
                self.page["ops"].append(
                    f"0.95 g {self.MARGIN:.2f} {self.y - lh:.2f} {self.content_w:.2f} {lh:.2f} re f 0 g")
                if chunk:
                    self.emit_line(self.MARGIN + 6, self.y - size, [(chunk, "c", None)], size)
                self.y -= lh
        self.y -= 8

    def rule(self):
        self.ensure(12)
        self.y -= 6
        self.page["ops"].append(
            f"0.7 G 0.5 w {self.MARGIN:.2f} {self.y:.2f} m "
            f"{self.W - self.MARGIN:.2f} {self.y:.2f} l S 0 G")
        self.y -= 8

    def table(self, rows, size=9.5):
        ncols = max(len(r) for r in rows)
        colw, pad = self.content_w / ncols, 4.0
        lh = size * 1.3
        for r_index, row in enumerate(rows):
            cells = [self.wrap(self.segments(row[c] if c < len(row) else "", r_index == 0),
                               size, colw - 2 * pad) for c in range(ncols)]
            height = max(len(c) for c in cells) * lh + 2 * pad
            self.ensure(height)
            top = self.y
            for c, lines in enumerate(cells):
                x = self.MARGIN + c * colw
                if r_index == 0:
                    self.page["ops"].append(
                        f"0.92 g {x:.2f} {top - height:.2f} {colw:.2f} {height:.2f} re f 0 g")
                self.page["ops"].append(
                    f"0.5 w 0.6 G {x:.2f} {top - height:.2f} {colw:.2f} {height:.2f} re S 0 G")
                for n, line in enumerate(lines):
                    if line:
                        self.emit_line(x + pad, top - pad - size - n * lh, line, size)
            self.y -= height
        self.y -= 10

    # -- keluaran -----------------------------------------------------------
    def render(self, title=None):
        objects = []

        def add(body):
            objects.append(body if isinstance(body, bytes) else body.encode("latin-1"))
            return len(objects)

        add("<< /Type /Catalog /Pages 2 0 R >>")
        add("PAGES_PLACEHOLDER")
        font_refs = []
        for key in ("r", "b", "i", "bi", "c"):
            num = add(f"<< /Type /Font /Subtype /Type1 /BaseFont /{_PDF_FONTS[key][1]} "
                      f"/Encoding /WinAnsiEncoding >>")
            font_refs.append(f"/{_PDF_FONTS[key][0]} {num} 0 R")
        fonts = " ".join(font_refs)
        info = "<< /Producer (MDs " + VERSION + ")"
        if title:
            info += " /Title <FEFF" + title.encode("utf-16-be").hex().upper() + ">"
        info += " >>"
        info_num = add(info)

        page_nums = []
        for page in self.pages:
            data = zlib.compress("\n".join(page["ops"]).encode("latin-1", "replace"))
            content = add(b"<< /Length %d /Filter /FlateDecode >>\nstream\n" % len(data)
                          + data + b"\nendstream")
            annots = []
            for x1, y1, x2, y2, url in page["links"]:
                safe = re.sub(r"[^\x21-\x7e]", lambda m: "%%%02X" % ord(m.group(0)), url)
                safe = safe.replace("\\", "%5C").replace("(", "%28").replace(")", "%29")
                annots.append(str(add(
                    f"<< /Type /Annot /Subtype /Link /Rect [{x1:.2f} {y1:.2f} {x2:.2f} {y2:.2f}] "
                    f"/Border [0 0 0] /A << /S /URI /URI ({safe}) >> >>")) + " 0 R")
            annot_txt = f" /Annots [{' '.join(annots)}]" if annots else ""
            page_nums.append(add(
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {self.W} {self.H}] "
                f"/Resources << /Font << {fonts} >> >> /Contents {content} 0 R{annot_txt} >>"))
        kids = " ".join(f"{n} 0 R" for n in page_nums)
        objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_nums)} >>".encode("latin-1")

        out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
        offsets = []
        for num, body in enumerate(objects, 1):
            offsets.append(len(out))
            out += b"%d 0 obj\n" % num + body + b"\nendobj\n"
        xref = len(out)
        out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
        for off in offsets:
            out += b"%010d 00000 n \n" % off
        out += (b"trailer\n<< /Size %d /Root 1 0 R /Info %d 0 R >>\nstartxref\n%d\n%%%%EOF\n"
                % (len(objects) + 1, info_num, xref))
        return bytes(out)


def md_to_pdf(md_text, out_path, title=None):
    """Bangun file .pdf dari teks Markdown. Return (diganti_tanda_tanya, emoji_dihapus)."""
    pdf = _PdfBuilder()
    if PAGE_MARK_RE.search(md_text):  # hasil ekstraksi PDF: teks apa adanya
        for kind, data in _md_blocks_plain(md_text):
            if kind == "pagebreak":
                pdf.page_break()
                continue
            for line in data.split("\n"):
                pdf.paragraph([(t, "r", None) for t in [_to_winansi(line, pdf.lost)]],
                              size=10, after=0, lead=1.35)
            pdf.y -= 8
    else:
        counters = {}
        for kind, data in _md_blocks(md_text):
            if kind != "li":
                counters.clear()
            if kind == "pagebreak":
                pdf.page_break()
            elif kind == "h":
                level, text = data
                size = {1: 20, 2: 16, 3: 14, 4: 12.5, 5: 11.5}.get(level, 11)
                if not pdf.at_top():
                    pdf.y -= 8
                pdf.ensure(size * 1.4 + 40)
                pdf.paragraph(pdf.segments(text, True), size=size, after=6, lead=1.25)
            elif kind == "p":
                pdf.paragraph(pdf.segments(_join_paragraph(data, False)))
            elif kind == "li":
                level, ordered, text = data
                for deeper in [k for k in counters if k > level]:
                    del counters[deeper]
                if ordered:
                    counters[level] = counters.get(level, 0) + 1
                    marker = f"{counters[level]}."
                else:
                    counters.pop(level, None)
                    marker = "\u2022"
                pdf.paragraph(pdf.segments(text), indent=22 + 18 * level, marker=marker,
                              marker_x=6 + 18 * level, after=2)
            elif kind == "code":
                pdf.code(data)
            elif kind == "quote":
                pdf.paragraph(pdf.segments(_join_paragraph(data, False)), indent=16, bar=True)
            elif kind == "hr":
                pdf.rule()
            elif kind == "table":
                pdf.table(data)
    parent = os.path.dirname(out_path)
    if parent:
        os.makedirs(parent, exist_ok=True)
    with open(out_path, "wb") as f:
        f.write(pdf.render(title))
    return tuple(pdf.lost)


# --------------------------------------------------------------------------
# Inti: File -> Markdown
# --------------------------------------------------------------------------

EMBED_TAG = "<!-- META_ORIGINAL_BASE64"
MAX_EMBED_BYTES = 25_000_000


def make_embed_block(raw):
    """Blok komentar HTML berisi salinan file asli (base64) agar bisa dipulihkan persis."""
    if len(raw) > MAX_EMBED_BYTES:
        raise ValueError(f"file terlalu besar untuk disematkan ({format_size(len(raw))})")
    b64 = base64.b64encode(raw).decode("ascii")
    wrapped = "\n".join(b64[i:i + 76] for i in range(0, len(b64), 76))
    return f"\n{EMBED_TAG} sha256={hashlib.sha256(raw).hexdigest()}\n{wrapped}\n-->\n"


def extract_embedded(text):
    """Pisahkan blok salinan asli. Return (teks_tanpa_blok, bytes | None, rusak)."""
    start = text.rfind(EMBED_TAG)
    if start < 0:
        return text, None, False
    end = text.find("-->", start)
    if end < 0:
        return text, None, True
    head, _, payload = text[start + len(EMBED_TAG):end].partition("\n")
    body = text[:start].rstrip("\n") + "\n" + text[end + 3:].lstrip("\n")
    body = body.rstrip("\n") + "\n" if body.strip() else ""
    match = re.search(r"sha256=([0-9a-fA-F]{64})", head)
    try:
        raw = base64.b64decode("".join(payload.split()), validate=True)
    except Exception:
        return body, None, True
    if not match or hashlib.sha256(raw).hexdigest() != match.group(1).lower():
        return body, None, True
    return body, raw, False


def convert_file_to_md(path, embed=False):
    ext = os.path.splitext(path)[1].lower()
    name = os.path.basename(path)
    if ext == ".pdf":
        body = pdf_to_text(path)
    elif ext == ".docx":
        body = docx_to_markdown(path)
    elif ext == ".doc":
        raise ValueError("format .doc lama tidak didukung; simpan ulang sebagai .docx")
    elif lang_for(name):
        body = read_text_file(path)
    else:
        raise ValueError(f"format '{ext or name}' tidak didukung")
    if embed and ext in DOC_EXTS:
        with open(path, "rb") as f:
            body = body.rstrip("\n") + "\n" + make_embed_block(f.read())
    return f"<!-- META_NAME: {name} -->\n\n{body}"


META_RE = re.compile(r"<!--\s*META_NAME:\s*(.+?)\s*-->")


def split_meta(content):
    """Pisahkan header META_NAME. Return (nama_asli | None, isi)."""
    match = META_RE.match(content)
    if not match:
        return None, content
    rest = content[match.end():]
    if rest.startswith("\n\n"):
        rest = rest[2:]
    elif rest.startswith("\n"):
        rest = rest[1:]
    return sanitize_filename(match.group(1)) or None, rest


def write_restored(target, body, original=None, orig_ext=None, rebuild=False):
    """Tulis hasil pemulihan satu file. Return (mode, catatan).

    Ekstensi `target` menentukan format: .docx / .pdf dibangun dari Markdown
    (atau disalin persis dari salinan asli bila ada), lainnya ditulis sebagai teks.
    """
    ext = os.path.splitext(target)[1].lower()
    parent = os.path.dirname(target)
    if parent:
        os.makedirs(parent, exist_ok=True)
    stem = os.path.splitext(os.path.basename(target))[0]
    if ext in DOC_EXTS and original is not None and not rebuild and ext == orig_ext:
        with open(target, "wb") as f:
            f.write(original)
        return "asli", "dipulihkan persis dari salinan asli yang tertanam"
    if ext == ".docx":
        md_to_docx(body, target, title=stem)
        return "docx", "dibangun ulang dari teks Markdown; tata letak & gambar asli tidak ikut"
    if ext == ".pdf":
        replaced, dropped = md_to_pdf(body, target, title=stem)
        note = "dibangun ulang dari teks Markdown; tata letak & gambar asli tidak ikut"
        if replaced:
            note += f"; {replaced} karakter non-Latin diganti '?'"
        if dropped:
            note += f"; {dropped} emoji/simbol dihapus"
        return "pdf", note
    write_text_file(target, body)
    return "teks", ""


def clean_path(raw):
    """Bersihkan path hasil seret-lepas / tempel: kutip, 'file://', '& ' PowerShell, spasi ter-escape."""
    p = raw.strip()
    if p.startswith("& "):
        p = p[2:].strip()
    if len(p) >= 2 and p[0] == p[-1] and p[0] in "\"'":
        p = p[1:-1]
    if p.lower().startswith("file://"):
        p = unquote(urlparse(p).path)
        if os.name == "nt" and re.match(r"^/[A-Za-z]:", p):
            p = p[1:]
    elif os.name != "nt":
        p = re.sub(r"\\(.)", r"\1", p)
    return os.path.expanduser(p)


# --------------------------------------------------------------------------
# Inti: Folder -> Markdown (bundle)
# --------------------------------------------------------------------------

def make_fence(text):
    """Pagar code block yang lebih panjang dari backtick mana pun di dalam teks."""
    longest = max((len(m) for m in re.findall(r"`{3,}", text)), default=2)
    return "`" * max(3, longest + 1)


def build_tree(rel_paths, root_name):
    tree = {}
    for rel in rel_paths:
        node = tree
        for part in rel.split("/"):
            node = node.setdefault(part, {})

    lines = [root_name + "/"]

    def walk(node, prefix):
        items = sorted(node.items(), key=lambda kv: (not kv[1], kv[0].lower()))
        for index, (name, child) in enumerate(items):
            last = index == len(items) - 1
            lines.append(prefix + ("└── " if last else "├── ") + name + ("/" if child else ""))
            if child:
                walk(child, prefix + ("    " if last else "│   "))

    walk(tree, "")
    return "\n".join(lines)


def is_prior_bundle(path):
    if not path.lower().endswith(".md"):
        return False
    try:
        with open(path, "rb") as f:
            return f.read(64).startswith(BUNDLE_TITLE.encode("utf-8"))
    except OSError:
        return False


def parse_ignore(values):
    """Ubah teks/daftar pola (dipisah koma) menjadi list pola."""
    if not values:
        return []
    if isinstance(values, str):
        values = [values]
    patterns = []
    for value in values:
        for part in value.split(","):
            part = part.strip().replace("\\", "/").strip("/")
            if part:
                patterns.append(part)
    return patterns


def is_ignored(rel_path, patterns):
    """Cocokkan path relatif (pemisah '/') dengan pola gaya shell (fnmatch).

    Pola cocok bila mengenai path lengkap, nama file, atau salah satu nama folder.
    Contoh: 'tests', '*.min.js', 'docs/*', 'src/legacy'.
    """
    if not patterns:
        return False
    parts = rel_path.split("/")
    prefixes = ["/".join(parts[:i]) for i in range(1, len(parts) + 1)]
    for pat in patterns:
        if any(fnmatch.fnmatch(x, pat) for x in prefixes) or \
                any(fnmatch.fnmatch(p, pat) for p in parts):
            return True
    return False


def scan_folder(target_dir, exts=None, max_bytes=MAX_FILE_BYTES, ignore=None, exclude_path=None):
    """Pindai folder tanpa menulis apa pun.

    Return (included [(rel, bahasa, teks)], skipped [(rel, alasan)], jumlah_diabaikan).
    """
    ignore = ignore or []
    exclude_abs = os.path.abspath(exclude_path) if exclude_path else None
    included, skipped, ignored = [], [], 0

    for root, dirs, files in os.walk(target_dir):
        dirs[:] = sorted(d for d in dirs if d not in IGNORE_DIRS and not d.startswith("."))

        for name in sorted(files):
            low = name.lower()
            if (name.startswith(".") and low not in SAFE_DOTFILES) or name in IGNORE_FILES:
                continue
            lang = lang_for(name)
            if lang is None:
                continue
            ext = os.path.splitext(low)[1]
            if exts and not ({ext, "." + low, low} & exts):
                continue
            path = os.path.join(root, name)
            rel = os.path.relpath(path, target_dir).replace(os.sep, "/")
            if os.path.islink(path):
                real = os.path.realpath(path)
                base = os.path.realpath(target_dir)
                try:
                    outside = os.path.commonpath([base, real]) != base
                except ValueError:  # beda drive di Windows
                    outside = True
                if outside:
                    skipped.append((rel, "symlink ke luar folder"))
                    continue
            if is_ignored(rel, ignore):
                ignored += 1
                continue
            if os.path.abspath(path) == exclude_abs or is_prior_bundle(path):
                continue
            try:
                size = os.path.getsize(path)
                if size > max_bytes:
                    skipped.append((rel, f"terlalu besar: {format_size(size)}"))
                    continue
                included.append((rel, lang, read_text_file(path)))
            except ValueError as e:
                skipped.append((rel, str(e)))
            except OSError as e:
                skipped.append((rel, f"gagal dibaca: {e}"))
    return included, skipped, ignored


def bundle_folder(target_dir, output_md, exts=None, max_bytes=MAX_FILE_BYTES,
                  ignore=None, scanned=None):
    """Gabungkan file teks/kode dalam folder ke satu file Markdown.

    `scanned` (hasil scan_folder) boleh diberikan agar folder tidak dipindai dua kali.
    Return dict: count, skipped, ignored, chars, tokens, text.
    """
    if scanned is None:
        scanned = scan_folder(target_dir, exts, max_bytes, ignore, exclude_path=output_md)
    included, skipped, ignored = scanned
    folder_name = os.path.basename(os.path.abspath(target_dir)) or "project_bundle"

    if not included:
        return {"count": 0, "skipped": skipped, "ignored": ignored,
                "chars": 0, "tokens": 0, "text": ""}

    tree = build_tree([rel for rel, _, _ in included], folder_name)
    tree_fence = make_fence(tree)
    parts = [
        f"{BUNDLE_TITLE} {folder_name}\n",
        f"> Dibuat oleh MDs v{VERSION} pada {datetime.now():%Y-%m-%d %H:%M} "
        f"| {len(included)} file\n",
        f"## Struktur Direktori\n\n{tree_fence}text\n{tree}\n{tree_fence}\n\n---\n",
    ]
    for rel, lang, text in included:
        fence = make_fence(text)
        parts.append(f"## File: `{rel}`\n\n{fence}{lang}\n{text}\n{fence}\n\n---\n")

    result = "\n".join(parts)
    write_text_file(output_md, result)
    return {
        "count": len(included),
        "skipped": skipped,
        "ignored": ignored,
        "chars": len(result),
        "tokens": estimate_tokens(result),
        "text": result,
    }


# --------------------------------------------------------------------------
# Inti: Markdown -> File asli (unbundle)
# --------------------------------------------------------------------------

FILE_HDR_RE = re.compile(r"^## File: `(.+)`\s*$")
OPEN_FENCE_RE = re.compile(r"^(`{3,})[^`]*$")


def parse_bundle(content):
    """Return daftar [(rel_path, isi_file)] dari file bundle."""
    lines = content.split("\n")
    entries = []
    i = 0
    while i < len(lines):
        header = FILE_HDR_RE.match(lines[i])
        if header:
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            fence_match = OPEN_FENCE_RE.match(lines[j]) if j < len(lines) else None
            if fence_match:
                fence = fence_match.group(1)
                k = j + 1
                while k < len(lines) and lines[k].rstrip() != fence:
                    k += 1
                if k < len(lines):
                    entries.append((header.group(1), "\n".join(lines[j + 1:k])))
                    i = k + 1
                    continue
        i += 1
    return entries


def extract_entries(entries, dest_dir, overwrite=False):
    """Tulis semua entri bundle. Return [(rel, status, pesan)]."""
    results = []
    for rel, body in entries:
        target = safe_join(dest_dir, rel)
        if target is None:
            results.append((rel, "ditolak", "path tidak aman"))
        elif os.path.exists(target) and not overwrite:
            results.append((rel, "dilewati", "sudah ada"))
        else:
            try:
                write_text_file(target, body)
                results.append((rel, "ok", ""))
            except OSError as e:
                results.append((rel, "gagal", str(e)))
    return results


# --------------------------------------------------------------------------
# Navigasi folder interaktif
# --------------------------------------------------------------------------

def browse_and_select(current_dir, target_type="file", allowed_exts=None, accept_named=False):
    """Pilih file atau folder lewat nomor. Return path, atau None bila batal."""
    page = 0
    while True:
        current_dir = os.path.abspath(current_dir)
        try:
            names = sorted(os.listdir(current_dir), key=str.lower)
        except OSError as e:
            print(f"{RED}Error membaca direktori: {e}{RESET}")
            parent = os.path.dirname(current_dir)
            if parent == current_dir:
                return None
            current_dir, page = parent, 0
            pause()
            continue

        folders = [n for n in names
                   if not n.startswith(".") and os.path.isdir(os.path.join(current_dir, n))]
        files = []
        if target_type == "file":
            files = [n for n in names
                     if not n.startswith(".")
                     and os.path.isfile(os.path.join(current_dir, n))
                     and (allowed_exts is None
                          or os.path.splitext(n)[1].lower() in allowed_exts
                          or (accept_named and lang_for(n) is not None))]
        entries = [(n, True) for n in folders] + [(n, False) for n in files]
        total_pages = max(1, -(-len(entries) // PAGE_SIZE))
        page = min(page, total_pages - 1)
        select_keys = {"s"} | ({"99"} if len(entries) < 99 else set())

        print_banner()
        print(f"{BLUE}[Direktori Saat Ini]: {current_dir}{RESET}")
        print("=" * 50)
        print(f"0. {YELLOW}[Kembali ke direktori sebelumnya]{RESET}")
        if target_type == "folder":
            print(f"s. {GREEN}[PILIH DIREKTORI INI]{RESET}")

        start = page * PAGE_SIZE
        for number, (name, is_dir) in enumerate(entries[start:start + PAGE_SIZE], start + 1):
            if is_dir:
                print(f"{number}. {CYAN}[DIR] {name}{RESET}")
            else:
                print(f"{number}. {name} {MAGENTA}({os.path.splitext(name)[1].lower()}){RESET}")
        if not entries:
            print("(kosong)")
        if total_pages > 1:
            print(f"-- Halaman {page + 1}/{total_pages}  (n = berikutnya, p = sebelumnya) --")
        print("-" * 50)

        raw = ask("Pilih nomor (q = batal, exit = keluar) atau tempel/seret path: ")
        choice = raw.lower()

        if choice == "q":
            return None
        if choice == "n":
            page = min(page + 1, total_pages - 1)
            continue
        if choice == "p":
            page = max(page - 1, 0)
            continue
        if choice == "0":
            parent = os.path.dirname(current_dir)
            if parent == current_dir:
                print(f"{RED}Sudah berada di root direktori!{RESET}")
                pause()
            else:
                current_dir, page = parent, 0
            continue
        if target_type == "folder" and choice in select_keys:
            return current_dir

        try:
            number = int(choice)
        except ValueError:
            candidate = os.path.abspath(clean_path(raw)) if raw else ""
            if candidate and os.path.isdir(candidate):
                if target_type == "folder":
                    return candidate
                current_dir, page = candidate, 0
            elif candidate and os.path.isfile(candidate):
                if target_type == "folder":
                    print(f"{RED}Itu file, bukan folder.{RESET}")
                    pause()
                elif allowed_exts is not None \
                        and os.path.splitext(candidate)[1].lower() not in allowed_exts \
                        and not (accept_named and lang_for(os.path.basename(candidate))):
                    print(f"{RED}Format file ini tidak didukung untuk menu ini.{RESET}")
                    pause()
                else:
                    return candidate
            else:
                print(f"{RED}Masukkan nomor atau path yang valid!{RESET}")
                pause()
            continue
        if 1 <= number <= len(entries):
            name, is_dir = entries[number - 1]
            if is_dir:
                current_dir, page = os.path.join(current_dir, name), 0
            else:
                return os.path.join(current_dir, name)
        else:
            print(f"{RED}Pilihan tidak valid!{RESET}")
            pause()


def choose_destination(default_dir):
    """Pilih lokasi simpan hasil ekstraksi. Return path, atau None bila batal."""
    print("\nLokasi penyimpanan hasil:")
    print(f"1. Direktori saat ini ({default_dir})")
    print("2. Pilih direktori lain")
    print("3. Buat folder baru di direktori saat ini")
    choice = ask("Pilihan [1]: ") or "1"

    if choice == "2":
        return browse_and_select(default_dir, target_type="folder")
    if choice == "3":
        name = ask("Nama folder baru: ")
        if not name or sanitize_filename(name) != name:
            print(f"{RED}Nama folder tidak valid.{RESET}")
            return None
        path = os.path.join(default_dir, name)
        try:
            os.makedirs(path, exist_ok=True)
        except OSError as e:
            print(f"{RED}Gagal membuat folder: {e}{RESET}")
            return None
        return path
    return default_dir


# --------------------------------------------------------------------------
# Menu interaktif
# --------------------------------------------------------------------------

def print_bundle_preview(included, skipped, ignored=0, limit=30):
    """Tampilkan daftar file yang akan dibundle. limit=None -> tampilkan semua."""
    total_chars = sum(len(text) for _, _, text in included)
    print(f"\n{GREEN}{len(included)} file akan digabungkan{RESET} "
          f"(~{format_size(total_chars)}, ~{max(1, total_chars // 4):,} token)")
    shown = included if limit is None else included[:limit]
    for rel, _, text in shown:
        print(f"  + {rel} ({format_size(len(text.encode('utf-8')))})")
    if len(shown) < len(included):
        print(f"  ... dan {len(included) - len(shown)} file lainnya")
    if ignored:
        print(f"{CYAN}{ignored} file diabaikan oleh pola abaikan{RESET}")
    if skipped:
        print(f"{YELLOW}{len(skipped)} file dilewati:{RESET}")
        shown_skipped = skipped if limit is None else skipped[:10]
        for rel, reason in shown_skipped:
            print(f"  - {rel}: {reason}")
        if len(shown_skipped) < len(skipped):
            print(f"  ... dan {len(skipped) - len(shown_skipped)} lainnya")


def print_restore_preview(entries, dest_dir, limit=30):
    """Tampilkan daftar file yang akan diekstrak. Return hitungan per status."""
    rows = []
    counts = {"baru": 0, "sudah ada": 0, "ditolak": 0}
    for rel, _ in entries:
        target = safe_join(dest_dir, rel)
        if target is None:
            status = "ditolak"
        elif os.path.exists(target):
            status = "sudah ada"
        else:
            status = "baru"
        counts[status] += 1
        rows.append((rel, status))

    print(f"\n{GREEN}{len(entries)} file akan diekstrak ke:{RESET} {dest_dir}")
    print(f"  baru: {counts['baru']} | sudah ada: {counts['sudah ada']} "
          f"| ditolak: {counts['ditolak']}")
    style = {
        "baru": (GREEN, "+", ""),
        "sudah ada": (YELLOW, "~", " (sudah ada)"),
        "ditolak": (RED, "x", " (path tidak aman)"),
    }
    shown = rows if limit is None else rows[:limit]
    for rel, status in shown:
        color, mark, note = style[status]
        print(f"  {color}{mark} {rel}{note}{RESET}")
    if len(shown) < len(rows):
        print(f"  ... dan {len(rows) - len(shown)} file lainnya")
    return counts


def confirm_preview(prompt, printer, limit=30):
    """Tampilkan pratinjau lalu minta konfirmasi. 'l' = lihat daftar lengkap."""
    printer(limit)
    while True:
        answer = ask(f"{prompt} (y/n, l = lihat semua): ").lower()
        if answer in ("l", "lihat"):
            printer(None)
            continue
        return answer in ("y", "ya")


END_MENU = "\nTekan [Enter] untuk kembali ke menu utama..."
END_DONE = "\nTekan [Enter] untuk selesai..."


def convert_file_flow(selected, end_msg=END_MENU):
    """Alur konversi satu file ke Markdown (dipakai menu 1 dan seret-lepas)."""
    ext = os.path.splitext(selected)[1].lower()
    name = os.path.basename(selected)
    if ext not in DOC_EXTS and ext != ".doc" and lang_for(name) is None:
        print(f"{RED}Format '{ext or name}' belum didukung.{RESET}")
        pause(end_msg)
        return

    default_out = os.path.splitext(selected)[0] + ".md"
    print_banner()
    print(f"File terpilih: {selected}")
    custom = ask(f"Nama file hasil Markdown [{os.path.basename(default_out)}]: ")
    output = default_out
    if custom:
        safe = sanitize_filename(custom)
        if safe:
            if not safe.lower().endswith(".md"):
                safe += ".md"
            output = os.path.join(os.path.dirname(selected), safe)

    if os.path.abspath(output) == os.path.abspath(selected):
        stem = os.path.splitext(selected)[0]
        output = f"{stem}_converted.md"
        print(f"{YELLOW}Nama sama dengan file sumber, hasil disimpan sebagai "
              f"'{os.path.basename(output)}'.{RESET}")

    embed = False
    if ext in DOC_EXTS:
        print(f"{CYAN}Markdown hanya berisi teks. Salinan file asli bisa disematkan agar "
              f"nanti dipulihkan persis (ukuran file .md bertambah ~33%).{RESET}")
        embed = confirm(f"Sematkan salinan asli {ext[1:].upper()}?")

    if os.path.exists(output) and not confirm(f"'{os.path.basename(output)}' sudah ada. Timpa?"):
        print("Dibatalkan.")
        pause(end_msg)
        return
    if not confirm(f"Konversi '{name}' menjadi '{os.path.basename(output)}'?"):
        print("Konversi dibatalkan.")
        pause(end_msg)
        return

    try:
        text = convert_file_to_md(selected, embed=embed)
        write_text_file(output, text)
        print(f"{GREEN}Berhasil! Tersimpan di: {output}{RESET}")
        print(f"Ukuran {format_size(len(text.encode('utf-8')))} | ~{estimate_tokens(text):,} token")
        if embed:
            print(f"{CYAN}Salinan asli tertanam: pulihkan lewat menu 3 / MDc restore.{RESET}")
    except Exception as e:
        print(f"{RED}Terjadi kesalahan: {e}{RESET}")
    pause(end_msg)


def menu_convert_file():
    print_banner()
    print("--- 1. Convert File ke Markdown ---")
    selected = browse_and_select(default_start_dir(), "file", SUPPORTED_EXTS | {".doc"},
                                 accept_named=True)
    if selected:
        convert_file_flow(selected)


def parse_exts(raw):
    if not raw:
        return None
    result = set()
    for part in re.split(r"[,\s]+", raw.strip().lower()):
        if part:
            result.add(part if part.startswith(".") else "." + part)
    return result or None


def bundle_flow(target_dir, end_msg=END_MENU):
    """Alur bundle folder ke satu Markdown (dipakai menu 2 dan seret-lepas)."""
    folder_name = os.path.basename(target_dir) or "project_bundle"
    default_name = f"{folder_name}.md"
    print_banner()
    print(f"Direktori terpilih: {target_dir}")
    custom = ask(f"Nama file bundle Markdown [{default_name}]: ")
    name = sanitize_filename(custom) if custom else default_name
    name = name or default_name
    if not name.lower().endswith(".md"):
        name += ".md"
    output = os.path.join(target_dir, name)

    exts = parse_exts(ask("Filter ekstensi (kosong = semua, contoh: .py,.js): "))
    ignore = parse_ignore(ask("Abaikan file/folder (kosong = tidak ada, contoh: tests,*.min.js): "))

    if os.path.exists(output) and not confirm(f"'{name}' sudah ada. Timpa?"):
        print("Dibatalkan.")
        pause(end_msg)
        return

    print("\nMemindai folder...")
    scanned = scan_folder(target_dir, exts, MAX_FILE_BYTES, ignore, exclude_path=output)
    included, skipped, ignored = scanned
    if not included:
        print_bundle_preview(included, skipped, ignored)
        print(f"{YELLOW}Tidak ada file yang cocok untuk digabungkan.{RESET}")
        pause(end_msg)
        return

    if not confirm_preview(
            f"Convert isi folder '{folder_name}' ke '{name}'?",
            lambda limit: print_bundle_preview(included, skipped, ignored, limit)):
        print("Konversi folder dibatalkan.")
        pause(end_msg)
        return

    try:
        stats = bundle_folder(target_dir, output, scanned=scanned)
    except OSError as e:
        print(f"{RED}Gagal membuat bundle: {e}{RESET}")
        pause(end_msg)
        return

    print(f"{GREEN}Berhasil menggabungkan {stats['count']} file ke '{name}'!{RESET}")
    print(f"Ukuran {format_size(stats['chars'])} | ~{stats['tokens']:,} token (perkiraan kasar)")

    if stats["count"] and shutil.which("termux-clipboard-set"):
        if confirm("Salin hasil ke clipboard?"):
            print(f"{GREEN}Tersalin!{RESET}" if copy_to_clipboard(stats["text"])
                  else f"{RED}Gagal menyalin (mungkin terlalu besar).{RESET}")
    pause(end_msg)


def menu_convert_folder():
    print_banner()
    print("--- 2. Convert Isi Folder ke Markdown (Bundle) ---")
    target_dir = browse_and_select(default_start_dir(), target_type="folder")
    if target_dir:
        bundle_flow(target_dir)


def restore_flow(selected, end_msg=END_MENU):
    """Alur pemulihan Markdown -> file asli / bundle (menu 3 dan seret-lepas)."""
    try:
        content = read_text_file(selected)
    except (OSError, ValueError) as e:
        print(f"{RED}Gagal membaca file: {e}{RESET}")
        pause(end_msg)
        return

    base_dir = os.path.dirname(selected)
    entries = parse_bundle(content)

    if entries:  # --- bundle proyek ---
        print_banner()
        print(f"{GREEN}[INFO] Terdeteksi Bundle Proyek: {len(entries)} file{RESET}")
        dest = choose_destination(base_dir)
        if not dest:
            print("Dibatalkan.")
            pause(end_msg)
            return
        counts = {}

        def show(limit):
            counts.update(print_restore_preview(entries, dest, limit))

        if not confirm_preview(f"Ekstrak {len(entries)} file ke '{dest}'?", show):
            print("Ekstraksi dibatalkan.")
            pause(end_msg)
            return
        overwrite = False
        if counts.get("sudah ada"):
            print(f"{YELLOW}{counts['sudah ada']} file sudah ada di tujuan.{RESET}")
            overwrite = confirm("Timpa file yang sudah ada? (n = lewati saja)")
        ok = 0
        for rel, status, msg in extract_entries(entries, dest, overwrite):
            if status == "ok":
                ok += 1
                print(f"{GREEN}[DIPULIHKAN] {rel}{RESET}")
            else:
                print(f"{RED}[{status.upper()}] {rel}: {msg}{RESET}")
        print(f"\n{GREEN}Berhasil mengekstrak {ok} dari {len(entries)} file ke {dest}{RESET}")
        pause(end_msg)
        return

    # --- file tunggal ---
    detected, rest = split_meta(content)
    body, original, broken = extract_embedded(rest)
    orig_ext = os.path.splitext(detected)[1].lower() if detected else None
    print_banner()
    print(f"File Markdown terpilih: {selected}")

    output_name = ""
    if detected:
        suggested = detected
        if orig_ext == ".doc":
            suggested = os.path.splitext(detected)[0] + ".docx"
        print(f"{GREEN}[INFO] Ditemukan nama asli: {detected}{RESET}")
        if broken:
            print(f"{RED}[PERINGATAN] Salinan asli tertanam rusak (checksum tidak cocok); "
                  f"dokumen dibangun ulang dari teks.{RESET}")
        if orig_ext in DOC_EXTS or orig_ext == ".doc":
            if original is not None and orig_ext in DOC_EXTS:
                print(f"{GREEN}[INFO] Salinan file asli tertanam: akan dipulihkan persis.{RESET}")
            else:
                print(f"{YELLOW}[INFO] Tanpa salinan asli, '{suggested}' dibangun ulang dari teks "
                      f"Markdown (judul, daftar, tabel, tebal/miring ikut; tata letak & gambar "
                      f"asli tidak).{RESET}")
        if confirm(f"Gunakan nama '{suggested}'?"):
            output_name = suggested
    else:
        print(f"{YELLOW}[INFO] Markdown biasa (tanpa metadata nama asli).{RESET}")
        print("Tips: beri nama berakhiran .docx atau .pdf untuk mengubahnya jadi dokumen.")
    if not output_name:
        output_name = sanitize_filename(
            ask("Nama file hasil (contoh: index.html / catatan.txt / laporan.docx / laporan.pdf): "))
    if not output_name:
        print("Nama file tidak valid / kosong.")
        pause(end_msg)
        return

    dest = choose_destination(base_dir)
    if not dest:
        print("Dibatalkan.")
        pause(end_msg)
        return
    output_path = os.path.join(dest, output_name)
    if os.path.exists(output_path) and not confirm(f"'{output_path}' sudah ada. Timpa?"):
        print("Dibatalkan.")
        pause(end_msg)
        return
    if not confirm(f"Simpan ke '{output_path}'?"):
        print("Dibatalkan.")
        pause(end_msg)
        return
    try:
        mode, note = write_restored(output_path, body, original, orig_ext)
        print(f"{GREEN}Berhasil memulihkan file ke: {output_path}{RESET}")
        if note:
            print(f"{CYAN}Catatan: {note}{RESET}")
    except Exception as e:
        print(f"{RED}Terjadi kesalahan: {e}{RESET}")
    pause(end_msg)


def menu_restore():
    print_banner()
    print("--- 3. Convert Markdown ke File Asli / Unbundle ---")
    selected = browse_and_select(default_start_dir(), "file", MD_EXTS)
    if selected:
        restore_flow(selected)


def interactive_menu():
    while True:
        print_banner()
        print("Pilih Menu Utama:")
        print("1. Convert file -> Markdown")
        print("2. Convert isi folder -> Markdown (bundle)")
        print("3. Convert Markdown -> File asli (TXT/DOCX/PDF) / Unbundle")
        print("0. Keluar / ketik 'exit'")
        print(f"{CYAN}Tip: seret atau tempel path file/folder ke sini untuk langsung memprosesnya.{RESET}")
        missing = [n for n, lib in (("pypdf", pypdf), ("python-docx", DocxDocument)) if lib is None]
        if missing:
            print(f"{YELLOW}Belum terpasang: {', '.join(missing)} (PDF/DOCX terbatas). "
                  f"Pasang: python -m pip install {' '.join(missing)}{RESET}")
        choice = ask("\nMasukkan pilihan (0-3): ")
        if choice == "0":
            print("\nTerima kasih telah menggunakan MDs!")
            return
        elif choice == "1":
            menu_convert_file()
        elif choice == "2":
            menu_convert_folder()
        elif choice == "3":
            menu_restore()
        else:
            candidate = os.path.abspath(clean_path(choice)) if choice else ""
            if candidate and os.path.exists(candidate):
                run_auto([candidate])
            else:
                print(f"{RED}Pilihan tidak valid!{RESET}")
                pause("\nTekan [Enter] untuk mencoba lagi...")


# --------------------------------------------------------------------------
# Mode CLI
# --------------------------------------------------------------------------

def build_parser():
    parser = argparse.ArgumentParser(
        prog="MDc",
        description="MDs - Ultimate Markdown Converter (tanpa argumen = mode interaktif)")
    parser.add_argument("--version", action="version", version=f"MDs {VERSION}")
    sub = parser.add_subparsers(dest="command")

    p_file = sub.add_parser("file", help="konversi satu file ke Markdown")
    p_file.add_argument("path")
    p_file.add_argument("-o", "--output", help="file .md hasil (default: di samping sumber)")
    p_file.add_argument("--embed", action="store_true",
                        help="PDF/DOCX: sematkan salinan asli di dalam .md (bisa dipulihkan persis)")
    p_file.add_argument("--force", action="store_true", help="timpa file hasil bila sudah ada")

    p_bundle = sub.add_parser("bundle", help="gabungkan isi folder ke satu file Markdown")
    p_bundle.add_argument("folder")
    p_bundle.add_argument("-o", "--output", help="file .md hasil (default: <folder>/<nama>.md)")
    p_bundle.add_argument("--ext", help="filter ekstensi, mis. .py,.js")
    p_bundle.add_argument("--ignore", action="append", metavar="POLA",
                          help="abaikan file/folder sesuai pola, dipisah koma, boleh diulang "
                               "(mis. tests,*.min.js)")
    p_bundle.add_argument("--max-kb", type=int, default=MAX_FILE_BYTES // 1000,
                          help="lewati file teks lebih besar dari ini (KB)")
    p_bundle.add_argument("--force", action="store_true")
    p_bundle.add_argument("-n", "--dry-run", action="store_true",
                          help="tampilkan daftar file saja, tanpa menulis apa pun")

    sub.add_parser("doctor", help="periksa instalasi (Python, pustaka, PATH, git)")

    p_restore = sub.add_parser("restore", help="pulihkan file/bundle Markdown ke file asli")
    p_restore.add_argument("markdown")
    p_restore.add_argument("-d", "--dest", help="folder tujuan (default: folder file .md)")
    p_restore.add_argument("--name", help="nama file hasil, mis. laporan.docx / laporan.pdf "
                                          "(format ditentukan oleh ekstensinya)")
    p_restore.add_argument("--text", action="store_true",
                           help="simpan dokumen PDF/DOCX sebagai .txt (perilaku lama)")
    p_restore.add_argument("--rebuild", action="store_true",
                           help="abaikan salinan asli tertanam, bangun ulang dari teks Markdown")
    p_restore.add_argument("--force", action="store_true", help="timpa file yang sudah ada")
    p_restore.add_argument("-n", "--dry-run", action="store_true",
                           help="tampilkan daftar file saja, tanpa menulis apa pun")
    return parser


def run_doctor():
    """Periksa instalasi; membantu saat 'pip'/'MDc' tidak dikenali di Windows CMD."""
    import platform
    import sysconfig
    from importlib import metadata

    def version_of(dist):
        try:
            return metadata.version(dist)
        except Exception:
            return None

    def line(label, ok, detail):
        mark = f"{GREEN}OK   {RESET}" if ok else f"{RED}BELUM{RESET}"
        print(f"  [{mark}] {label:<12} {detail}")

    print(f"MDs v{VERSION} | Python {platform.python_version()} | "
          f"{platform.system()} {platform.release()}")
    print(f"Python di: {sys.executable}\n")
    ver = version_of("pypdf")
    line("pypdf", bool(pypdf), ver or "python -m pip install pypdf")
    ver = version_of("python-docx")
    line("python-docx", DocxDocument is not None, ver or "python -m pip install python-docx")
    git = shutil.which("git")
    line("git", bool(git), git or "opsional (untuk clone/update); bisa pakai unduhan ZIP")
    mdc = shutil.which("MDc")
    line("perintah MDc", bool(mdc), mdc or "tidak ada di PATH")
    if not mdc:
        user_key = "nt_user" if os.name == "nt" else "posix_user"
        folders = {sysconfig.get_path("scripts"), sysconfig.get_path("scripts", user_key)}
        print(f"\n{YELLOW}MDc belum bisa dipanggil langsung. Pilih salah satu:{RESET}")
        print("  1. Tanpa instal ulang : python md.py   atau   python -m md   (di folder proyek)")
        print("  2. Windows            : jalankan MDc.bat (di folder proyek), atau tambahkan salah satu")
        print("     folder berikut ke PATH lalu buka CMD baru:")
        for folder in sorted(f for f in folders if f):
            print(f"       {folder}")
    return 0


def run_cli(args):
    def fail(msg):
        print(f"Error: {msg}", file=sys.stderr)
        return 1

    if args.command == "doctor":
        return run_doctor()

    if args.command == "file":
        if not os.path.isfile(args.path):
            return fail(f"file tidak ditemukan: {args.path}")
        output = args.output or os.path.splitext(args.path)[0] + ".md"
        if os.path.abspath(output) == os.path.abspath(args.path):
            output = os.path.splitext(args.path)[0] + "_converted.md"
        if os.path.exists(output) and not args.force:
            return fail(f"'{output}' sudah ada (gunakan --force untuk menimpa)")
        try:
            text = convert_file_to_md(args.path, embed=args.embed)
            write_text_file(output, text)
        except Exception as e:
            return fail(e)
        print(f"OK: {output} (~{estimate_tokens(text):,} token)"
              + (" + salinan asli tertanam" if args.embed and "META_ORIGINAL" in text else ""))
        return 0

    if args.command == "bundle":
        if not os.path.isdir(args.folder):
            return fail(f"folder tidak ditemukan: {args.folder}")
        folder = os.path.abspath(args.folder)
        output = args.output or os.path.join(folder, os.path.basename(folder) + ".md")
        if os.path.exists(output) and not args.force and not args.dry_run:
            return fail(f"'{output}' sudah ada (gunakan --force untuk menimpa)")
        scanned = scan_folder(folder, parse_exts(args.ext), args.max_kb * 1000,
                              parse_ignore(args.ignore), exclude_path=output)
        included, skipped, ignored = scanned
        if not included:
            return fail("tidak ada file yang cocok untuk digabungkan")
        if args.dry_run:
            print_bundle_preview(included, skipped, ignored, limit=None)
            print("\n(dry-run: tidak ada file yang ditulis)")
            return 0
        try:
            stats = bundle_folder(folder, output, scanned=scanned)
        except OSError as e:
            return fail(e)
        print(f"OK: {stats['count']} file -> {output} (~{stats['tokens']:,} token)")
        if ignored:
            print(f"  diabaikan oleh pola: {ignored} file")
        for rel, reason in stats["skipped"]:
            print(f"  dilewati: {rel} ({reason})")
        return 0

    if args.command == "restore":
        if not os.path.isfile(args.markdown):
            return fail(f"file tidak ditemukan: {args.markdown}")
        try:
            content = read_text_file(args.markdown)
        except (OSError, ValueError) as e:
            return fail(e)
        dest = args.dest or os.path.dirname(os.path.abspath(args.markdown))
        entries = parse_bundle(content)
        if entries:
            if args.dry_run:
                print_restore_preview(entries, dest, limit=None)
                print("\n(dry-run: tidak ada file yang ditulis)")
                return 0
            bad = 0
            for rel, status, msg in extract_entries(entries, dest, args.force):
                print(f"  [{status}] {rel} {msg}".rstrip())
                bad += status not in ("ok", "dilewati")
            return 1 if bad else 0
        name, rest = split_meta(content)
        body, original, broken = extract_embedded(rest)
        orig_ext = os.path.splitext(name)[1].lower() if name else None
        if args.name:
            out_name = sanitize_filename(args.name)
        elif name:
            out_name = name
            if args.text and orig_ext in DOC_EXTS | {".doc"}:
                out_name = os.path.splitext(name)[0] + ".txt"
            elif orig_ext == ".doc":
                out_name = os.path.splitext(name)[0] + ".docx"
        else:
            return fail("bukan bundle dan tidak ada nama asli; gunakan --name "
                        "(mis. --name laporan.docx)")
        if not out_name:
            return fail("nama file tidak valid")
        target = safe_join(dest, out_name)
        if target is None:
            return fail("nama file tidak aman")
        if args.dry_run:
            print(f"Akan menulis: {target}" + (" (sudah ada)" if os.path.exists(target) else ""))
            print("\n(dry-run: tidak ada file yang ditulis)")
            return 0
        if os.path.exists(target) and not args.force:
            return fail(f"'{target}' sudah ada (gunakan --force untuk menimpa)")
        if broken:
            print("Peringatan: salinan asli tertanam rusak (checksum tidak cocok); "
                  "dibangun ulang dari teks.", file=sys.stderr)
        try:
            mode, note = write_restored(target, body, original, orig_ext, args.rebuild)
        except Exception as e:
            return fail(e)
        print(f"OK: {target}" + (f" ({note})" if note else ""))
        return 0

    return 0


COMMANDS = {"file", "bundle", "restore", "doctor"}


def run_auto(paths):
    """Seret-lepas / tempel path: aksi dipilih otomatis (folder -> bundle, .md -> pulihkan,
    file lain -> Markdown)."""
    for raw in paths:
        path = os.path.abspath(clean_path(raw))
        if os.path.isdir(path):
            bundle_flow(path, END_DONE)
        elif os.path.isfile(path):
            if os.path.splitext(path)[1].lower() in MD_EXTS:
                restore_flow(path, END_DONE)
            else:
                convert_file_flow(path, END_DONE)
        else:
            print(f"{RED}Tidak ditemukan: {raw}{RESET}")
            pause(END_DONE)
    return 0


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):  # cegah crash di konsol lama (non-UTF-8)
        try:
            stream.reconfigure(errors="replace")
        except Exception:
            pass
    argv = sys.argv[1:] if argv is None else list(argv)
    try:
        if argv and argv[0] not in COMMANDS and not argv[0].startswith("-"):
            return run_auto(argv)
        args = build_parser().parse_args(argv)
        if args.command is None:
            interactive_menu()
            return 0
        return run_cli(args)
    except (KeyboardInterrupt, EOFError):
        print("\nDibatalkan. Sampai jumpa!")
        return 130


if __name__ == "__main__":
    sys.exit(main())

